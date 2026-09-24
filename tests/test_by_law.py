from dataclasses import replace
from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
import pytest

from documents.by_law import SPEC
from document_generator import heading_map, preflight_template, validate_values, xml_parts
from models import RowError, SetupError
from processor import run_batch
from tests.extension_helpers import (ADDRESS, configured, render, register_for_test,
                                     assert_style_resources_unchanged)
from tests.helpers import FakeSheet

HEADERS = ['Association Name / संस्था का नाम', 'Association Address / संस्था का पंजीकृत पता',
           'Work Area / संस्था का कार्यक्षेत्र']
VALUES = dict(association_name='परीक्षण संस्था', association_address=ADDRESS, work_area='इंदौर जिला')


@pytest.fixture
def by_law_settings(settings, monkeypatch):
    register_for_test(monkeypatch, SPEC)
    return configured(settings, SPEC, 'By_Law_Template.docx', 'By-Law Responses')


def test_exact_contract_and_complete_address(by_law_settings):
    assert len(SPEC.fields) == 3 and all(f.required for f in SPEC.fields)
    assert heading_map(SPEC, HEADERS) == {key: i for i, key in enumerate(VALUES)}
    assert {f.placeholder for f in SPEC.fields} == set(VALUES)
    assert not {'khasra_number', 'project_location'} & set(VALUES)
    config = by_law_settings.documents[SPEC.key]
    preflight_template(SPEC, config, by_law_settings)
    context, output = render(SPEC, by_law_settings, VALUES)
    assert context == VALUES
    text = xml_parts(BytesIO(output))['word/document.xml']
    for value in VALUES.values():
        assert value in text
    assert text.count(ADDRESS) == 1
    assert not any(tag in text for tag in ('{{', '{%', '}}', '%}'))
    assert_style_resources_unchanged(config.template_path, output)


@pytest.mark.parametrize('parameter', list(VALUES))
@pytest.mark.parametrize('blank', ['', ' \n\t ', None])
def test_required_fields(by_law_settings, parameter, blank):
    with pytest.raises(RowError, match=parameter):
        validate_values(SPEC, {**VALUES, parameter: blank}, by_law_settings.documents[SPEC.key], by_law_settings)


@pytest.mark.parametrize('change', ['unknown', 'missing'])
def test_template_contract(by_law_settings, tmp_path, change):
    config = by_law_settings.documents[SPEC.key]
    doc = Document(config.template_path)
    if change == 'unknown':
        doc.add_paragraph('{{unknown}}')
    else:
        doc.tables[0].cell(2, 2).text = ''
    path = tmp_path / 'bad.docx'
    doc.save(path)
    with pytest.raises(SetupError, match='placeholder mismatch'):
        preflight_template(SPEC, replace(config, template_path=path), by_law_settings)


def test_template_layout_is_preserved(by_law_settings):
    # The supplied full By-Law is NOT a one-page template. Preserve its explicit
    # break, paragraph/run formatting and table geometry instead of shrinking it.
    config = by_law_settings.documents[SPEC.key]
    _, output = render(SPEC, by_law_settings, VALUES)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    with ZipFile(config.template_path) as source, ZipFile(BytesIO(output)) as generated:
        before = ET.fromstring(source.read('word/document.xml'))
        after = ET.fromstring(generated.read('word/document.xml'))
        for path in ('.//w:pPr', './/w:rPr', './/w:tblPr', './/w:tblGrid', './/w:tcPr', './/w:sectPr', './/w:br'):
            assert [ET.canonicalize(ET.tostring(x)) for x in before.findall(path, ns)] == [
                ET.canonicalize(ET.tostring(x)) for x in after.findall(path, ns)]
        assert len(before.findall('.//w:br[@w:type="page"]', ns)) >= 1
        assert len(Document(BytesIO(output)).paragraphs) == len(Document(config.template_path).paragraphs)


def test_dry_run_is_readonly(by_law_settings):
    sheet = FakeSheet(SPEC.key, [VALUES])
    original = [row.copy() for row in sheet.data]
    result = run_batch(by_law_settings, SPEC.key, rows=[2], dry_run=True, worksheet=sheet)
    assert result.counts['would_generate'] == 1
    assert sheet.data == original and sheet.writes == []
    assert not by_law_settings.documents[SPEC.key].output_dir.exists()


def test_duplicate_or_missing_heading():
    with pytest.raises(SetupError, match='Ambiguous'):
        heading_map(SPEC, HEADERS + HEADERS[:1])
    with pytest.raises(SetupError, match='work_area'):
        heading_map(SPEC, HEADERS[:-1])
