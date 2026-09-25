from dataclasses import replace
from io import BytesIO
import json
from docx import Document
import pytest
from document_registry import SPECS
from document_generator import heading_map, preflight_template, xml_parts
from models import RowError, SetupError
from tests.helpers import values
from tests.extension_helpers import ROOT, render

SPEC = SPECS['consent']


def member_values(count):
    data = values('consent')
    for i in range(1, 13):
        for part in ('name', 'designation'):
            if i <= count:
                data[f'member_{i}_{part}'] = f'श्री परीक्षण {i} पिता श्री पूर्ण नाम' if part == 'name' else f'पद {i}'
            elif i <= 11:
                data[f'member_{i}_{part}'] = ''
    return data


@pytest.mark.parametrize('count', [5, 6, 11])
def test_dynamic_rows_and_address_semantics(settings, count):
    data = member_values(count)
    data['project_location'] = 'स्वतंत्र स्थान भोपाल'
    assert not {'society_address', 'survey_numbers', 'society_name'} & data.keys()
    preflight_template(SPEC, settings.documents['consent'], settings)
    context, output = render(SPEC, settings, data)
    parts = xml_parts(BytesIO(output))
    assert data['project_location'] in parts['word/header1.xml']
    assert data['association_name'] in parts['word/header1.xml']
    assert data['association_address'] in parts['word/document.xml']
    assert data['association_address'] not in parts['word/header1.xml']
    for key, value in context.items():
        if key != 'members':
            assert value in ''.join(parts.values())
    assert len(context['members']) == count
    doc = Document(BytesIO(output))
    assert len(doc.tables) == 1
    assert len(doc.tables[0].rows) == count + 1
    for i,row in enumerate(doc.tables[0].rows[1:],1):
        assert [c.text for c in row.cells] == [f'{i:02d}', data[f'member_{i}_name'], data[f'member_{i}_designation']]
    assert not any(tag in ''.join(parts.values()) for tag in ('{{', '}}', '{%', '%}'))


def test_actual_headers_and_pasted_newline_normalization():
    headers = json.loads((ROOT / 'tests/fixtures/migration_headers.json').read_text(encoding='utf-8'))['consent']
    mapping = heading_map(SPEC, headers)
    assert len(mapping) == 35
    assert mapping['association_name'] == 1 and mapping['member_9_designation'] == 31
    # Retain coverage for the original live newlines after the user cleaned them.
    headers[1] = '\n' + headers[1]
    headers[31] = '\n' + headers[31]
    assert heading_map(SPEC, headers) == mapping
    headers[1] = '"' + headers[1] + '"'
    with pytest.raises(SetupError, match='association_name'):
        heading_map(SPEC, headers)  # Literal quotes are not guessed away.


@pytest.mark.parametrize('count', [4, 12])
def test_member_count_boundaries(settings, count):
    with pytest.raises(RowError, match='member'):
        render(SPEC, settings, member_values(count))


@pytest.mark.parametrize('i', range(6,12))
@pytest.mark.parametrize('part', ['name','designation'])
def test_partial_optional_member(settings, i, part):
    data = member_values(i-1)
    data[f'member_{i}_{part}'] = 'supplied'
    with pytest.raises(RowError, match=f'member_{i}'):
        render(SPEC, settings, data)


def test_optional_gap(settings):
    data = member_values(7)
    data['member_6_name'] = data['member_6_designation'] = ' \n '
    with pytest.raises(RowError, match='contiguous'):
        render(SPEC, settings, data)


@pytest.mark.parametrize('field', [f.parameter for f in SPEC.fields if f.required])
def test_missing_required_input(settings, field):
    data = member_values(5)
    del data[field]
    with pytest.raises(RowError, match=field):
        render(SPEC, settings, data)


@pytest.mark.parametrize('change', ['missing','loop','unknown'])
def test_template_contract(settings, tmp_path, change):
    config = settings.documents['consent']
    doc = Document(config.template_path)
    if change == 'loop':
        doc.tables[0].cell(1,0).text = '{%tr for member in members[:5] %}'
    elif change == 'missing':
        doc.tables[0].cell(2,2).text = ''
    else:
        doc.add_paragraph('{{member.email}}')
    path = tmp_path / 'bad.docx'
    doc.save(path)
    with pytest.raises(SetupError):
        preflight_template(SPEC, replace(config, template_path=path), settings)


def test_ambiguous_and_excess_member_headers():
    headers = [f.heading for f in SPEC.fields]
    with pytest.raises(SetupError, match='Ambiguous'):
        heading_map(SPEC, headers + [headers[0]])
    with pytest.raises(SetupError, match='Unsupported member'):
        heading_map(SPEC, headers + ['Member 12 Name / सदस्य 12 का नाम'])
