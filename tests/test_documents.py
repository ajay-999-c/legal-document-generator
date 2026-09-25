from dataclasses import replace
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile
import re
import pytest
from docx import Document
from document_registry import SPECS
from document_generator import (validate_values, render_bytes, xml_parts, preflight_template,
    heading_map, safe_component, output_filename, save_atomic)
from models import RowError, SetupError
from tests.helpers import values
from tests.conftest import ROOT

# These fixtures encode the original three Forms (including their optional dates).
# New document contracts and formatting are covered in their own test modules.
LEGACY_KEYS = ('noc', 'affidavit', 'consent')


def render(key, settings, input_values=None):
    spec = SPECS[key]
    data = validate_values(spec, values(key) if input_values is None else input_values, settings.documents[key], settings)
    context = spec.context_builder(data)
    return data, context, render_bytes(settings.documents[key].template_path, context)


def noc_title_paragraph(document):
    paragraphs = list(document.paragraphs)
    for section in document.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            paragraphs.extend(header.paragraphs)
    return next(p for p in paragraphs if '{{association_name}}' in p.text)


@pytest.mark.parametrize('key,count,required', [('noc',12,11),('affidavit',14,14),('consent',35,22)])
def test_schema_against_business_contract(key,count,required):
    from tests.final_contract import assert_contract
    spec = SPECS[key]
    assert len(spec.fields) == count
    assert sum(f.required for f in spec.fields) == required
    assert_contract(key)


@pytest.mark.parametrize('key', LEGACY_KEYS)
def test_actual_template_render_text_and_format(settings,key):
    preflight_template(SPECS[key], settings.documents[key], settings)
    data, context, output=render(key,settings)
    parts=xml_parts(BytesIO(output))
    all_text=''.join(parts.values())
    for field in SPECS[key].fields:
        assert data[field.parameter] in all_text
    assert '{{' not in all_text
    if any(f.parameter == 'document_date' for f in SPECS[key].fields):
        assert '..........' in all_text
    # Static package styling survives; existing paragraphs/runs are not rebuilt.
    with ZipFile(settings.documents[key].template_path) as source, ZipFile(BytesIO(output)) as generated:
        for name in source.namelist():
            if name.startswith(('word/styles','word/numbering','word/theme','word/media')):
                if name.endswith('.xml'):
                    # Serialization may normalize declaration quotes/newlines.
                    # Compare XML content for every adapter, binary assets exactly.
                    assert ET.canonicalize(source.read(name)) == ET.canonicalize(generated.read(name))
                else:
                    assert source.read(name)==generated.read(name)
    if key=='affidavit':
        body=parts['word/document.xml']
        assert body.count(data['association_address']) == 1
        assert 'ASSOCIATION_NAME' in context and 'association_address' in context
        assert 'कार्यकारिणी कोषाध्यक्ष' in body and body.count('_________') == 3
        assert not {'LAND', 'PROJECT_LOCATION', 'DESIGNATION', 'DATE'} & context.keys()
    if key=='consent':
        assert data['project_location'] in parts['word/header1.xml']
        assert data['association_name'] in parts['word/header1.xml']
        positions=[parts['word/document.xml'].index(data[f'member_{i}_name']) for i in range(1,6)]
        assert positions==sorted(positions)
        assert len(context['members'])==5


@pytest.mark.parametrize('key,param', [(key,f.parameter) for key in LEGACY_KEYS for f in SPECS[key].fields if f.required])
def test_every_required_field_rejected(settings,key,param):
    data=values(key); data[param]=' \n\t '
    with pytest.raises(RowError,match=param):
        validate_values(SPECS[key],data,settings.documents[key],settings)


@pytest.mark.parametrize('association', [' परीक्षण संघ ', 'परीक्षण संघ – भोपाल', 'परीक्षण संघ – इंदौर'])
def test_noc_city_is_never_rewritten(settings,association):
    data=values('noc'); data['association_name']=association
    validated,context,output=render('noc',settings,data)
    assert context['association_name']==association.strip()
    paragraphs = xml_parts(BytesIO(output))
    assert association.strip() in [text.strip() for key, text in paragraphs.items()
                                   if key.startswith('__paragraph__')]


@pytest.mark.parametrize('key', LEGACY_KEYS)
def test_identifiers_xml_and_multiline(settings,key):
    data=values(key)
    field=next(f.parameter for f in SPECS[key].fields if f.kind=='text')
    data[field]='  परीक्षण & <TEST>\nदूसरी पंक्ति  '
    validated,ctx,output=render(key,settings,data)
    assert validated[field]=='परीक्षण & <TEST>\nदूसरी पंक्ति'
    assert 'परीक्षण & <TEST>' in ''.join(xml_parts(BytesIO(output)).values())
    if 'plot_no' in data: assert validated['plot_no']=='026'


@pytest.mark.parametrize('fmt,value,expected', [('%Y-%m-%d','2026-09-23','23-09-2026'),('%m/%d/%Y','03/04/2026','04-03-2026'),('%d/%m/%Y','03/04/2026','03-04-2026')])
def test_explicit_date_conventions(settings,fmt,value,expected):
    config=replace(settings.documents['noc'],input_date_format=fmt)
    data=values('noc'); data['completion_certificate_date']=value
    assert validate_values(SPECS['noc'],data,config,settings)['completion_certificate_date']==expected


@pytest.mark.parametrize('value',['03/04/2026','2026-02-30','today','..........'])
def test_invalid_or_unconfigured_date(settings,value):
    data=values('noc'); data['completion_certificate_date']=value
    with pytest.raises(RowError,match='completion_certificate_date'):
        validate_values(SPECS['noc'],data,settings.documents['noc'],settings)


@pytest.mark.parametrize('age',['0','121','20.5','True','-1','1e2'])
def test_age_range(settings,age):
    data=values('affidavit'); data['age']=age
    with pytest.raises(RowError,match='age'):
        validate_values(SPECS['affidavit'],data,settings.documents['affidavit'],settings)


def test_heading_normalization_aliases_and_ambiguity():
    spec=SPECS['noc']; labels=[f.heading for f in spec.fields]
    changed=['\ufeff  '+s.upper().replace(' / ','/\n ')+'   ' for s in labels]
    assert heading_map(spec,changed)['association_name']==0
    assert heading_map(spec,changed)['project_location']==1
    with pytest.raises(SetupError,match='Ambiguous'):
        heading_map(spec,labels+[labels[2]])
    with pytest.raises(SetupError,match='signatory_name'):
        heading_map(spec,labels[:-1])
    changed[1]='Association Address / एसोसिएशन का पता'
    with pytest.raises(SetupError,match='project_location'):
        heading_map(spec,changed)
    aff=SPECS['affidavit']; labels=[f.heading for f in aff.fields]
    labels[1]="Father's Name / पिता का नाम"
    assert heading_map(aff,labels)['father_name']==1
    with pytest.raises(SetupError, match='Ambiguous'):
        heading_map(aff, labels + [aff.fields[1].heading])
    labels[5]='Unreviewed Address'
    with pytest.raises(SetupError,match='association_address'):
        heading_map(aff,labels)


@pytest.mark.parametrize('token',['{{unknown}}','{{ members[5].name }}'])
def test_unexpected_template_and_nested_path(settings,tmp_path,token):
    path=tmp_path/'bad.docx'; doc=Document(settings.documents['consent'].template_path)
    doc.add_paragraph(token); doc.save(path)
    with pytest.raises(SetupError,match='mismatch'):
        preflight_template(SPECS['consent'],replace(settings.documents['consent'],template_path=path),settings)


def test_strict_nested_render(settings):
    spec=SPECS['consent']; data=values('consent'); ctx=spec.context_builder(data)
    del ctx['members'][2]['designation']
    with pytest.raises(RowError): render_bytes(settings.documents['consent'].template_path,ctx)


def test_split_run_placeholder_preflight(settings,tmp_path):
    source=settings.documents['noc'].template_path
    doc=Document(source); paragraph=noc_title_paragraph(doc); paragraph.clear()
    paragraph.add_run('{{association_'); paragraph.add_run('name}}')
    path=tmp_path/'split.docx'; doc.save(path)
    preflight_template(SPECS['noc'],replace(settings.documents['noc'],template_path=path),settings)


def test_safe_filename_and_atomic_save(settings,tmp_path,monkeypatch):
    assert safe_component('../../CON\x00').find('/')==-1
    assert safe_component('CON')=='_CON'
    data=values('noc'); data['project_name']='परीक्षण'*200
    config=settings.documents['noc']; name=output_filename(config,data,42)
    assert len(name.encode())<=230 and name.endswith('.docx')
    target=tmp_path/name; target.write_bytes(b'old')
    def fail(*args): raise OSError('save failure')
    monkeypatch.setattr(Path,'replace',fail)
    with pytest.raises(RowError): save_atomic(b'new',target)
    assert target.read_bytes()==b'old'
    assert not list(tmp_path.glob('.docx-*'))

def test_missing_optional_header_and_duplicate_member():
    spec=SPECS['consent']; headers=[f.heading for f in spec.fields]
    with pytest.raises(SetupError,match='document_date'):
        heading_map(spec,[h for h in headers if h!='Document Date / दस्तावेज़ दिनांक'])
    with pytest.raises(SetupError,match='Ambiguous'):
        heading_map(spec,headers+[spec.fields[11].heading])


def test_missing_placeholder_and_old_noc_title(settings,tmp_path):
    config=settings.documents['noc']
    doc=Document(config.template_path); title=noc_title_paragraph(doc); title.add_run(' – इंदौर')
    path=tmp_path/'old.docx'; doc.save(path)
    with pytest.raises(SetupError,match='title'):
        preflight_template(SPECS['noc'],replace(config,template_path=path),settings)
    title.text='fixed title'; doc.save(path)
    with pytest.raises(SetupError,match='mismatch'):
        preflight_template(SPECS['noc'],replace(config,template_path=path),settings)


def test_footer_unexpected_placeholder(settings,tmp_path):
    config=settings.documents['affidavit']; doc=Document(config.template_path)
    doc.sections[0].footer.paragraphs[0].text='{{UNSUPPORTED_FOOTER}}'
    path=tmp_path/'footer.docx'; doc.save(path)
    with pytest.raises(SetupError,match='mismatch'):
        preflight_template(SPECS['affidavit'],replace(config,template_path=path),settings)


def make_noc_layout(tmp_path, title_location='header', suffix='', missing=None):
    """Synthetic DOCX only: never writes to the supplied templates."""
    doc = Document()
    title = '{{association_name}}' + suffix
    if title_location == 'body':
        doc.add_paragraph('दिनांक: {{document_date}}')
        doc.add_paragraph(title)
        doc.add_paragraph('{{project_location}}')
    else:
        container = getattr(doc.sections[0], title_location)
        container.paragraphs[0].text = title
        container.add_paragraph('{{project_location}}')
        doc.add_paragraph('दिनांक: {{document_date}}')
    for field in SPECS['noc'].fields:
        if field.parameter not in {'association_name', 'project_location', 'document_date'}:
            doc.add_paragraph('{{' + field.placeholder + '}}')
    if missing:
        containers = [doc, doc.sections[0].header, doc.sections[0].first_page_header,
                      doc.sections[0].even_page_header, doc.sections[0].footer]
        for container in containers:
            for paragraph in container.paragraphs:
                if '{{' + missing + '}}' in paragraph.text:
                    paragraph.text = paragraph.text.replace('{{' + missing + '}}', '')
    path = tmp_path / 'synthetic_noc.docx'
    doc.save(path)
    return path


@pytest.mark.parametrize('title_location', ['header', 'first_page_header', 'even_page_header', 'body'])
def test_noc_title_in_header_or_after_body_date(settings, tmp_path, title_location):
    path = make_noc_layout(tmp_path, title_location)
    config = replace(settings.documents['noc'], template_path=path)
    preflight_template(SPECS['noc'], config, settings)
    doc = Document(path)
    assert next(p.text for p in doc.paragraphs if p.text.strip()) == 'दिनांक: {{document_date}}'
    data = validate_values(SPECS['noc'], values('noc'), config, settings)
    output = render_bytes(path, SPECS['noc'].context_builder(data))
    parts = xml_parts(BytesIO(output))
    location_text = parts['word/document.xml'] if title_location == 'body' else ''.join(
        text for name, text in parts.items() if re.fullmatch(r'word/header\d+\.xml', name))
    assert data['association_name'] in location_text
    assert data['project_location'] in location_text


@pytest.mark.parametrize('title_location', ['header', 'body'])
@pytest.mark.parametrize('suffix', [' – इंदौर', ' - इंदौर', ' — इंदौर', ' इंदौर', ' – भोपाल', '\n– इंदौर'])
def test_noc_fixed_city_still_rejected(settings, tmp_path, title_location, suffix):
    path = make_noc_layout(tmp_path, title_location, suffix)
    with pytest.raises(SetupError, match='NOC title'):
        preflight_template(SPECS['noc'], replace(settings.documents['noc'], template_path=path), settings)


@pytest.mark.parametrize('missing', ['association_name', 'project_location'])
def test_noc_missing_header_placeholder_still_rejected(settings, tmp_path, missing):
    path = make_noc_layout(tmp_path, missing=missing)
    with pytest.raises(SetupError, match='placeholder mismatch'):
        preflight_template(SPECS['noc'], replace(settings.documents['noc'], template_path=path), settings)


def test_noc_valid_header_does_not_hide_bad_body_title(settings, tmp_path):
    path = make_noc_layout(tmp_path)
    doc = Document(path)
    doc.add_paragraph('{{association_name}} – इंदौर')
    doc.save(path)
    with pytest.raises(SetupError, match='NOC title'):
        preflight_template(SPECS['noc'], replace(settings.documents['noc'], template_path=path), settings)


def test_noc_footer_only_title_not_accepted(settings, tmp_path):
    path = make_noc_layout(tmp_path, 'footer')
    with pytest.raises(SetupError, match='NOC title'):
        preflight_template(SPECS['noc'], replace(settings.documents['noc'], template_path=path), settings)
