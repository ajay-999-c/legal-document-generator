from dataclasses import replace
from io import BytesIO
import json
import subprocess
import sys

from docx import Document
import pytest

from documents.registration import SPEC
from document_generator import heading_map, output_filename, preflight_template, validate_values, xml_parts
from models import RowError, SetupError
from processor import run_batch
from tests.extension_helpers import (ADDRESS, ROOT, configured, render,
                                     register_for_test, assert_style_resources_unchanged)
from tests.helpers import FakeSheet


HEADERS = ['Authority Location / सक्षम प्राधिकारी का स्थान ', 'Project Name / परियोजना का नाम ',
           'Association Address', 'Police Station / पुलिस थाना ',
           'Association / Society Name / एसोसिएशन / संस्था का नाम ', 'Place / स्थान ',
           'Signatory Name / हस्ताक्षरकर्ता का नाम ']
VALUES = dict(authority_location='राऊ', project_name='परीक्षण विहार', association_address=ADDRESS,
              police_station='राऊ थाना', association_name='परीक्षण संस्था', place='इंदौर',
              signatory_name='श्री परीक्षण सिंह')


@pytest.fixture
def registration_settings(settings, monkeypatch):
    register_for_test(monkeypatch, SPEC)
    return configured(settings, SPEC, 'Registration_Template.docx', 'Registration Responses')


def test_real_contract_and_complete_address(registration_settings):
    assert len(SPEC.fields) == 7 and all(f.required for f in SPEC.fields)
    assert heading_map(SPEC, HEADERS) == {key: i for i, key in enumerate(VALUES)}
    assert {f.placeholder for f in SPEC.fields} == set(VALUES)
    assert not {'khasra_number', 'project_location', 'document_date'} & set(VALUES)
    config = registration_settings.documents[SPEC.key]
    preflight_template(SPEC, config, registration_settings)
    context, output = render(SPEC, registration_settings, VALUES)
    text = xml_parts(BytesIO(output))['word/document.xml']
    assert context == VALUES
    for value in VALUES.values():
        assert value in text
    assert text.count(ADDRESS) == 1
    assert not any(tag in text for tag in ('{{', '{%', '}}', '%}'))
    assert_style_resources_unchanged(config.template_path, output)


@pytest.mark.parametrize('parameter', list(VALUES))
@pytest.mark.parametrize('missing', [False, True])
def test_required_fields(registration_settings, parameter, missing):
    data = dict(VALUES)
    if missing:
        del data[parameter]
    else:
        data[parameter] = ' \n\t '
    with pytest.raises(RowError, match=parameter):
        validate_values(SPEC, data, registration_settings.documents[SPEC.key], registration_settings)


@pytest.mark.parametrize('change', ['unknown', 'missing'])
def test_template_contract(registration_settings, tmp_path, change):
    config = registration_settings.documents[SPEC.key]
    doc = Document(config.template_path)
    if change == 'unknown':
        doc.add_paragraph('{{unknown}}')
    else:
        for paragraph in doc.paragraphs:
            if '{{association_address}}' in paragraph.text:
                paragraph.text = paragraph.text.replace('{{association_address}}', '')
    path = tmp_path / 'bad.docx'
    doc.save(path)
    with pytest.raises(SetupError, match='placeholder mismatch'):
        preflight_template(SPEC, replace(config, template_path=path), registration_settings)


def test_headers_are_explicit():
    with pytest.raises(SetupError, match='association_address'):
        heading_map(SPEC, [h for h in HEADERS if h != 'Association Address'])
    with pytest.raises(SetupError, match='Ambiguous'):
        heading_map(SPEC, HEADERS + [' Association Address '])


def test_filename_sanitization(registration_settings):
    data = {**VALUES, 'association_name': '../../CON\\ bad:<>|?*'}
    name = output_filename(registration_settings.documents[SPEC.key], data, 12)
    assert name.endswith('_12.docx')
    assert not any(c in name for c in '/\\:<>|?*') and '..' not in name


def test_selected_dry_run(registration_settings, tmp_path):
    sheet = FakeSheet(SPEC.key, [VALUES, VALUES])
    original = [row.copy() for row in sheet.data]
    result = run_batch(registration_settings, SPEC.key, rows=[3], dry_run=True, worksheet=sheet)
    assert result.counts['would_generate'] == 1
    assert [r.row_number for r in result.rows] == [3]
    assert sheet.data == original and sheet.writes == []
    assert not registration_settings.documents[SPEC.key].output_dir.exists()
    assert not list(tmp_path.rglob('*.docx'))


def test_source_from_other_working_directory(tmp_path):
    # A fresh source interpreter exercises the same renderer without live credentials.
    code = '''
import sys
sys.path.insert(0, sys.argv[1])
from documents.registration import SPEC
from document_generator import render_bytes, xml_parts
from io import BytesIO
import json
from pathlib import Path
data=json.loads(sys.argv[2])
output=render_bytes(Path(sys.argv[1])/'templates/Registration_Template.docx', SPEC.context_builder(data))
assert data['association_address'] in xml_parts(BytesIO(output))['word/document.xml']
print('registration source OK')
'''
    result = subprocess.run([sys.executable, '-c', code, str(ROOT), json.dumps(VALUES)],
                            cwd=tmp_path, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert 'registration source OK' in result.stdout
