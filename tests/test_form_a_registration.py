from dataclasses import replace
from io import BytesIO
import re
from xml.etree import ElementTree as ET

from docx import Document
import pytest

from documents.form_a_registration import SPEC
from document_generator import (heading_map, preflight_template, render_bytes,
                                validate_values, xml_parts)
from models import RowError, SetupError
from processor import run_batch
from tests.extension_helpers import (ADDRESS, configured, render, register_for_test,
                                     assert_style_resources_unchanged)
from tests.helpers import FakeSheet

SCALAR_HEADERS = [
    'Association Name / प्रस्तावित संस्था का नाम ', 'Association Address / संस्था का पता ',
    'Work Area / कार्यक्षेत्र ', 'Share Capital / अंश पूंजी ', 'Price Per Share / प्रति अंश कीमत ',
    'Number of Members / सदस्यों की संख्या ', 'Association Email / संस्था ईमेल ',
    'Project Name / परियोजना का नाम ', 'Completion Certificate Number / पूर्णता प्रमाण-पत्र क्रमांक ',
    'Completion Certificate Date / पूर्णता प्रमाण-पत्र दिनांक ', 'District / जिला ',
    'Management Committee Address / प्रबंध कार्यकारिणी का पता ',
    'Meeting Chairperson Name / बैठक अध्यक्ष का नाम ', 'Proposed By / प्रस्ताव रखने वाले का नाम ',
    'Approved By / अनुमोदन करने वाले का नाम ',
]
SCALARS = dict(association_name='परीक्षण संस्था', association_address=ADDRESS, work_area='धार कार्यक्षेत्र',
               share_capital='123456', price_per_share='700', member_count='85', association_email='',
               project_name='परीक्षण परियोजना', completion_certificate_no='00026/343',
               completion_certificate_date='2026-09-23', district_name='इंदौर',
               management_committee_address='परीक्षण कार्यालय राऊ',
               meeting_chairperson_name='स्वतंत्र अध्यक्ष', proposed_by='स्वतंत्र प्रस्तावक',
               approved_by='स्वतंत्र अनुमोदक')
MEMBER_KEYS = ('name', 'designation', 'plot_no', 'mobile')


def values(count=5, member_count=85):
    data = {**SCALARS, 'member_count': str(member_count)}
    for i in range(1, 12):
        member = dict(name=f'श्री परीक्षण {i} पिता श्री पूर्ण नाम {i}',
                      designation='अध्यक्ष' if i == 1 else f'पद {i}',
                      plot_no=f'0{i}/A', mobile=f'090000000{i:02d}')
        for key in MEMBER_KEYS:
            data[f'committee_member_{i}_{key}'] = member[key] if i <= count else ''
    if count > 11:
        for key in MEMBER_KEYS:
            data[f'committee_member_12_{key}'] = 'unsupported'
    return data


@pytest.fixture
def form_a_settings(settings, monkeypatch):
    register_for_test(monkeypatch, SPEC)
    return configured(settings, SPEC, 'Form_A_Registration_Template.docx', 'Form A Registration Responses')


def test_field_contract():
    assert len(SPEC.fields) == 59 and sum(f.required for f in SPEC.fields) == 34
    headers = SCALAR_HEADERS.copy()
    for i in range(1, 12):
        headers.extend([f'Committee Member {i} Name / समिति सदस्य {i} का नाम ',
                        f'Committee Member {i} Designation / समिति सदस्य {i} का पद ',
                        f'Committee Member {i} Plot Number / समिति सदस्य {i} भूखंड क्रमांक ',
                        f'Committee Member {i} Mobile Number / समिति सदस्य {i} मोबाइल नंबर '])
    assert heading_map(SPEC, headers) == {key: i for i, key in enumerate(values())}
    assert {f.parameter for f in SPEC.fields} == set(values())
    assert not {'meeting_date', 'khasra_number', 'project_location', 'father_name'} & set(values())
    assert [f.parameter for f in SPEC.fields if 'email' in f.parameter] == ['association_email']
    for i in range(1, 12):
        fields = [f for f in SPEC.fields if f.parameter.startswith(f'committee_member_{i}_')]
        assert len(fields) == 4 and all(f.required == (i <= 5) for f in fields)
        assert {f.placeholder for f in fields} == {f'committee_members[{i-1}].{k}' for k in MEMBER_KEYS}


@pytest.mark.parametrize('committee_count,member_count', [(5, 85), (6, 20), (11, 85), (5, 5)])
def test_actual_dynamic_tables(form_a_settings, committee_count, member_count):
    data = values(committee_count, member_count)
    config = form_a_settings.documents[SPEC.key]
    preflight_template(SPEC, config, form_a_settings)
    context, output = render(SPEC, form_a_settings, data)
    assert len(context['committee_members']) == committee_count
    assert context['member_count'] == member_count
    assert all(set(m) == set(MEMBER_KEYS) for m in context['committee_members'])
    assert context['share_capital'] == '123456' and context['price_per_share'] == '700'
    assert context['association_email'] == ''
    doc = Document(BytesIO(output))
    assert len(doc.tables) == 5
    # Every repeated table uses exactly the same ordered committee values.
    for table in doc.tables[1:4]:
        assert len(table.rows) - 1 == committee_count
        for i, row in enumerate(table.rows[1:], 1):
            assert row.cells[1].text == data[f'committee_member_{i}_name']
    for table in doc.tables[1:3]:
        for i, row in enumerate(table.rows[1:], 1):
            assert row.cells[0].text == f'{i:02d}'
            assert row.cells[2].text == data[f'committee_member_{i}_designation']
            assert row.cells[3].text == ''
    for i, row in enumerate(doc.tables[3].rows[1:], 1):
        assert row.cells[2].text == data[f'committee_member_{i}_plot_no']
        assert row.cells[3].text == data[f'committee_member_{i}_mobile']
    final_rows = doc.tables[-1].rows[1:]
    assert len(final_rows) == member_count  # Header is not a member row.
    for i, row in enumerate(final_rows, 1):
        expected = [str(i), '', 'सदस्य', '', '']
        if i <= committee_count:
            expected[1:4] = [data[f'committee_member_{i}_{k}'] for k in ('name', 'designation', 'plot_no')]
        assert [c.text for c in row.cells] == expected
    assert doc.tables[0].cell(7, 2).text == data['committee_member_1_name']
    assert doc.tables[0].cell(8, 2).text == data['committee_member_1_mobile']
    assert doc.tables[0].cell(9, 2).text == ''
    body = '\n'.join(p.text for p in doc.paragraphs)
    assert re.findall(r'दिनांक (\.+) को', body) == ['.' * 13]
    assert data['committee_member_1_name'] in body
    assert ADDRESS in body and doc.tables[0].cell(1, 2).text == ADDRESS
    for key in ('meeting_chairperson_name', 'proposed_by', 'approved_by', 'work_area'):
        assert data[key] in body
    assert '23-09-2026' in body
    assert not any(tag in ''.join(xml_parts(BytesIO(output)).values()) for tag in ('{{', '}}', '{%', '%}'))
    assert_style_resources_unchanged(config.template_path, output)
    source = Document(config.template_path)
    for before, after in zip(source.tables, doc.tables):
        for tag in ('tblPr', 'tblGrid'):
            assert ET.canonicalize(getattr(before._tbl, tag).xml) == ET.canonicalize(getattr(after._tbl, tag).xml)
    for index in range(1, 5):
        for row in doc.tables[index].rows[1:]:
            for before, after in zip(source.tables[index].rows[2].cells, row.cells):
                assert ET.canonicalize(before._tc.tcPr.xml) == ET.canonicalize(after._tc.tcPr.xml)


@pytest.mark.parametrize('parameter', [f.parameter for f in SPEC.fields if f.required])
def test_every_required_field(form_a_settings, parameter):
    with pytest.raises(RowError, match=parameter):
        render(SPEC, form_a_settings, {**values(), parameter: ' \n\t '})


@pytest.mark.parametrize('count', [4, 12])
def test_invalid_committee_count(form_a_settings, count):
    with pytest.raises(RowError, match='committee'):
        render(SPEC, form_a_settings, values(count))


@pytest.mark.parametrize('index', range(6, 12))
@pytest.mark.parametrize('key', MEMBER_KEYS)
def test_partial_optional_member(form_a_settings, index, key):
    data = values(index - 1)
    data[f'committee_member_{index}_{key}'] = 'supplied'
    with pytest.raises(RowError, match=f'committee_member_{index}'):
        render(SPEC, form_a_settings, data)


def test_noncontiguous_optional_members_rejected(form_a_settings):
    data = values(7)
    for key in MEMBER_KEYS:
        data[f'committee_member_6_{key}'] = ''
    with pytest.raises(RowError, match='contiguous'):
        render(SPEC, form_a_settings, data)


@pytest.mark.parametrize('count', ['', '0', '-1', '4', '4.0', '5.5', 'True', 'five', '1e2', '85,0', 'NaN', 'Infinity'])
def test_invalid_member_count(form_a_settings, count):
    with pytest.raises(RowError, match='member_count'):
        render(SPEC, form_a_settings, values(5, count))


@pytest.mark.parametrize('count', ['5', '005', '5.0', '5.00'])
def test_integer_like_count(form_a_settings, count):
    context, _ = render(SPEC, form_a_settings, values(5, count))
    assert context['member_count'] == 5


def test_count_below_populated_committee(form_a_settings):
    with pytest.raises(RowError, match='member_count'):
        render(SPEC, form_a_settings, values(8, 7))


def test_roles_may_repeat_but_are_never_inferred(form_a_settings):
    data = values()
    for key in ('meeting_chairperson_name', 'proposed_by', 'approved_by'):
        data[key] = 'एक ही स्वतंत्र व्यक्ति'
    context, _ = render(SPEC, form_a_settings, data)
    assert context['meeting_chairperson_name'] == context['proposed_by'] == context['approved_by']
    assert context['meeting_chairperson_name'] != context['committee_members'][0]['name']


@pytest.mark.parametrize('change', ['unknown', 'missing', 'nested', 'loop', 'dots', 'footer'])
def test_bad_template_rejected(form_a_settings, tmp_path, change):
    config = form_a_settings.documents[SPEC.key]
    doc = Document(config.template_path)
    if change == 'unknown':
        doc.add_paragraph('{{unexpected}}')
    elif change == 'missing':
        doc.tables[0].cell(4, 2).text = ''
    elif change == 'nested':
        doc.add_paragraph('{{committee_members[0].email}}')
    elif change == 'loop':
        doc.tables[-1].cell(1, 0).text = '{%tr for i in range(member_count - 1) %}'
    elif change == 'dots':
        for p in doc.paragraphs:
            if '.' * 13 in p.text:
                p.text = p.text.replace('.' * 13, '.' * 12)
    else:
        doc.sections[0].footer.paragraphs[0].text = '{{unexpected}}'
    path = tmp_path / 'bad.docx'
    doc.save(path)
    with pytest.raises(SetupError):
        preflight_template(SPEC, replace(config, template_path=path), form_a_settings)


def test_duplicate_missing_and_unsupported_headers():
    headers = [f.heading for f in SPEC.fields]
    with pytest.raises(SetupError, match='Ambiguous'):
        heading_map(SPEC, headers + ['  Association Name / प्रस्तावित संस्था का नाम '])
    with pytest.raises(SetupError, match='Ambiguous'):
        heading_map(SPEC, headers + [headers[-1]])
    with pytest.raises(SetupError, match='association_email'):
        heading_map(SPEC, [h for h in headers if not h.startswith('Association Email')])
    with pytest.raises(SetupError, match='Unsupported committee'):
        heading_map(SPEC, headers + ['Committee Member 12 Name / समिति सदस्य 12 का नाम'])


def test_strict_nested_context(form_a_settings):
    config = form_a_settings.documents[SPEC.key]
    context = SPEC.context_builder(validate_values(SPEC, values(), config, form_a_settings))
    del context['committee_members'][0]['mobile']
    with pytest.raises(RowError):
        render_bytes(config.template_path, context)


def test_readonly_dry_run_valid_and_invalid(form_a_settings):
    data = values(6)
    data['committee_member_6_mobile'] = ''
    sheet = FakeSheet(SPEC.key, [values(), data])
    original = [row.copy() for row in sheet.data]
    result = run_batch(form_a_settings, SPEC.key, rows=[2, 3], dry_run=True, worksheet=sheet)
    assert result.counts['would_generate'] == 1 and result.counts['invalid'] == 1
    assert sheet.data == original and sheet.writes == []
    assert not form_a_settings.documents[SPEC.key].output_dir.exists()
